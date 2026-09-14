import json
import os
import base64
import uuid
import io

import psycopg2
import boto3
from openpyxl import load_workbook

RATING_TYPES = ('blitz', 'rapid')


def get_conn():
    return psycopg2.connect(os.environ['DATABASE_URL'], options=f"-c search_path={os.environ.get('MAIN_DB_SCHEMA', 'public')}")


def cors_headers():
    return {
        'Access-Control-Allow-Origin': '*',
        'Access-Control-Allow-Methods': 'GET, POST, OPTIONS',
        'Access-Control-Allow-Headers': 'Content-Type, X-Admin-Password',
    }


def is_admin(event):
    return event.get('headers', {}).get('X-Admin-Password') == os.environ.get('ADMIN_PASSWORD')


def get_s3():
    return boto3.client(
        's3',
        endpoint_url='https://bucket.poehali.dev',
        aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
        aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY'],
    )


def cdn_url(key):
    return f"https://cdn.poehali.dev/projects/{os.environ['AWS_ACCESS_KEY_ID']}/bucket/{key}"


def file_to_dict(row):
    return {
        'id': row[0],
        'rating_type': row[1],
        'file_name': row[2],
        'file_url': row[3],
        'total_rows': row[4],
        'matched_count': row[5],
        'uploaded_at': str(row[6]),
    }


def handler(event: dict, context) -> dict:
    """Загрузка Excel-файлов рейтинга ФШР (блиц/рапид), автообновление рейтингов пользователей по ID ФШР и история загрузок"""
    if event.get('httpMethod') == 'OPTIONS':
        return {'statusCode': 200, 'headers': {**cors_headers(), 'Access-Control-Max-Age': '86400'}, 'body': ''}

    method = event.get('httpMethod', 'GET')

    if not is_admin(event):
        return {'statusCode': 403, 'headers': cors_headers(), 'body': json.dumps({'error': 'Forbidden'})}

    conn = get_conn()
    cur = conn.cursor()

    if method == 'GET':
        cur.execute(
            "SELECT id, rating_type, file_name, file_url, total_rows, matched_count, uploaded_at "
            "FROM fsr_rating_files ORDER BY uploaded_at DESC LIMIT 100"
        )
        rows = [file_to_dict(r) for r in cur.fetchall()]
        cur.close()
        conn.close()
        return {'statusCode': 200, 'headers': cors_headers(), 'body': json.dumps({'files': rows})}

    body = json.loads(event.get('body') or '{}')
    action = body.get('_action', '')

    if method == 'POST' and action == 'upload':
        rating_type = body.get('rating_type')
        file_b64 = body.get('file_b64', '')
        file_name = body.get('file_name', 'rating.xlsx')

        if rating_type not in RATING_TYPES:
            cur.close(); conn.close()
            return {'statusCode': 400, 'headers': cors_headers(), 'body': json.dumps({'error': 'Некорректный тип рейтинга'})}
        if not file_b64:
            cur.close(); conn.close()
            return {'statusCode': 400, 'headers': cors_headers(), 'body': json.dumps({'error': 'Файл не передан'})}

        data = base64.b64decode(file_b64)

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.active
        except Exception:
            cur.close(); conn.close()
            return {'statusCode': 400, 'headers': cors_headers(), 'body': json.dumps({'error': 'Не удалось прочитать Excel-файл'})}

        column = 'fsr_rating_blitz' if rating_type == 'blitz' else 'fsr_rating_rapid'

        total_rows = 0
        matched_count = 0
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            fsr_id = str(row[0]).strip()
            if not fsr_id:
                continue
            try:
                rating_value = int(float(row[3])) if len(row) > 3 and row[3] is not None else None
            except (ValueError, TypeError):
                rating_value = None
            if rating_value is None:
                continue
            total_rows += 1
            cur.execute(f"UPDATE users SET {column} = %s WHERE fsr_id = %s", (rating_value, fsr_id))
            if cur.rowcount > 0:
                matched_count += cur.rowcount

        key = f"fsr-ratings/{rating_type}_{uuid.uuid4().hex[:12]}_{file_name}"
        s3 = get_s3()
        s3.put_object(Bucket='files', Key=key, Body=data, ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        file_url = cdn_url(key)

        cur.execute(
            "INSERT INTO fsr_rating_files (rating_type, file_name, file_url, total_rows, matched_count) "
            "VALUES (%s, %s, %s, %s, %s) RETURNING id, rating_type, file_name, file_url, total_rows, matched_count, uploaded_at",
            (rating_type, file_name, file_url, total_rows, matched_count)
        )
        new_row = cur.fetchone()
        conn.commit()
        cur.close()
        conn.close()
        return {'statusCode': 200, 'headers': cors_headers(), 'body': json.dumps({'file': file_to_dict(new_row)})}

    cur.close()
    conn.close()
    return {'statusCode': 400, 'headers': cors_headers(), 'body': json.dumps({'error': 'Bad request'})}
